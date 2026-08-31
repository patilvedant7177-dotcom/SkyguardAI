import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_frontend_document():
    doc = docx.Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    
    COLOR_PRIMARY = RGBColor(15, 32, 39)
    COLOR_ACCENT = RGBColor(0, 131, 176)
    COLOR_DARK = RGBColor(44, 62, 80)
    COLOR_GRAY = RGBColor(100, 110, 120)
    
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = COLOR_DARK
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # HEADER / TITLE
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(20)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("🎨 SkyguardAI Frontend & UI Architecture")
    run_title.font.name = 'Segoe UI'
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Comprehensive Technical Structure, Component Catalog & Design System")
    run_sub.font.name = 'Segoe UI'
    run_sub.font.size = Pt(13.5)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_ACCENT

    # Metadata banner table
    meta_table = doc.add_table(rows=1, cols=3)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    meta_data = [
        ("Framework & Language", "React 19 + TypeScript 5.8+"),
        ("Build Tooling", "Vite 6 + HashRouter"),
        ("Visualization Stack", "Leaflet + Recharts + Lucide")
    ]
    
    for i, (k, v) in enumerate(meta_data):
        cell = meta_table.rows[0].cells[i]
        set_cell_background(cell, "F0F4F8")
        set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r_k = p.add_run(f"{k}\n")
        r_k.font.size = Pt(8.5)
        r_k.font.bold = True
        r_k.font.color.rgb = COLOR_GRAY
        r_v = p.add_run(v)
        r_v.font.size = Pt(9.5)
        r_v.font.bold = True
        r_v.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI Semibold'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI Semibold'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = COLOR_ACCENT
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.font.bold = True
            r_b.font.color.rgb = COLOR_DARK
        p.add_run(text)
        return p

    def add_callout(text, title="NOTE"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        set_cell_background(cell, "EBF3FA")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        rt = p.add_run(f"[{title}] ")
        rt.font.bold = True
        rt.font.color.rgb = COLOR_ACCENT
        p.add_run(text)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_code_block(code_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        set_cell_background(cell, "F4F5F7")
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.0)
        run.font.color.rgb = RGBColor(30, 40, 50)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 1. OVERVIEW & TECH STACK
    # -------------------------------------------------------------
    add_h1("1. UI Architecture & Technology Stack")
    doc.add_paragraph(
        "The SkyguardAI frontend delivers a high-density, real-time command center for meteorological operations. "
        "Engineered with React 19, TypeScript, and Vite, it integrates live Server-Sent Events (SSE) streaming, "
        "interactive geospatial maps, responsive multi-area timeseries charts, and explainable AI feature attribution dashboards."
    )

    add_bullet("React 19 & TypeScript: Strongly-typed component architecture ensuring 100% compliance with backend API models.", "• ")
    add_bullet("React Router 6 (HashRouter): Client-side routing with deep-linkable views and zero server configuration hurdles.", "• ")
    add_bullet("Leaflet & React-Leaflet: Geospatial station mapping with CartoDB DarkMatter tiles and status-coded animated marker pins.", "• ")
    add_bullet("Recharts Visualization Suite: Responsive multi-variable telemetry charts with gradient fills, dual axes, and reference zones.", "• ")
    add_bullet("Lucide React: Crisp vector iconography for all status badges, trend meters, and telemetry channels.", "• ")
    add_bullet("Glassmorphism Design System: Dark theme with CSS backdrop blur, glowing neon badges, and subtle micro-interactions.", "• ")

    # -------------------------------------------------------------
    # 2. APPLICATION ROUTING & NAVIGATION
    # -------------------------------------------------------------
    add_h1("2. Application Routing & Navigation Layout")
    doc.add_paragraph("The application layout wraps all views with a persistent frosted navigation header:")

    route_table = doc.add_table(rows=1, cols=3)
    route_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    route_headers = ["Route URL", "Component", "Functional Scope"]
    for i, h in enumerate(route_headers):
        cell = route_table.rows[0].cells[i]
        set_cell_background(cell, "0F2027")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    route_rows = [
        ("/", "NetworkOverview.tsx", "Geospatial array overview, global station KPIs, status filters, and interactive station cards."),
        ("/alerts", "LiveAlerts.tsx", "Real-time SSE alert streaming listener, severity triage filters, and one-click acknowledgment."),
        ("/station/:id", "StationDetail.tsx", "High-resolution 72-hour multi-variable telemetry charts, min/max metrics, and metadata."),
        ("/why-flagged/:id", "WhyFlagged.tsx", "SHAP feature attribution charts, detector consensus breakdown, and meteorological narrative."),
        ("/health/:id", "SensorHealth.tsx", "Circular SVG health gauge (0-100), degradation trajectory, and RUL maintenance forecast."),
        ("/history", "History.tsx", "Searchable audit archive for acknowledged and resolved historical anomalies across all nodes.")
    ]

    for row in route_rows:
        r_cells = route_table.add_row().cells
        for idx, text in enumerate(row):
            set_cell_background(r_cells[idx], "FFFFFF" if len(route_table.rows) % 2 == 0 else "F9FAFC")
            set_cell_margins(r_cells[idx], top=80, bottom=80, left=100, right=100)
            p = r_cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.add_run(text).font.size = Pt(9.0)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------
    # 3. DETAILED PAGE CATALOG
    # -------------------------------------------------------------
    add_h1("3. Detailed Page Breakdown & Operational Workflows")

    add_h2("3.1 Network Overview (NetworkOverview.tsx)")
    doc.add_paragraph("Central command dashboard providing immediate situational awareness across the entire sensor topology:")
    add_bullet("Network Health KPI Ribbon: Total Monitored Stations, Normal (green), Degrading (amber), Fault (rose), and Offline (gray).", "1. ")
    add_bullet("Interactive Leaflet Geospatial Map: Renders all stations on DarkMatter tiles with pulsating status markers and informative popups.", "2. ")
    add_bullet("Filterable Station Grid: Displays GPS coordinates, altitude, data source (Real IMD vs Synthetic), and direct action links.", "3. ")

    add_h2("3.2 Live Telemetry Alerts (LiveAlerts.tsx)")
    doc.add_paragraph("Real-time anomaly triage workstation powered by Server-Sent Events (SSE):")
    add_bullet("Continuous SSE Stream Listener: Ingests new alerts in real-time from GET /alerts/stream and prepends them with zero page reload.", "1. ")
    add_bullet("Live Stream Pulse Indicator: Displays visual connection health (STREAM ACTIVE vs DISCONNECTED).", "2. ")
    add_bullet("Instant Search & Severity Filtering: Real-time search by summary or station name; severity chips (High, Medium, Low).", "3. ")
    add_bullet("Triage Card Actions: One-click alert acknowledgment (POST /alerts/{id}/acknowledge) and direct deep-link to XAI diagnostics.", "4. ")

    add_h2("3.3 Station Detail & Timeseries Telemetry (StationDetail.tsx)")
    doc.add_paragraph("Deep telemetry analysis interface for analyzing historical sensor behavior:")
    add_bullet("Station & Time Window Switcher: Select any station node; toggle time horizons (24h, 48h, 72h, 168h / 7 days).", "1. ")
    add_bullet("Live Parameter Metrics: Current Temperature (°C) with 72h min/max, Barometric Pressure (hPa), and Relative Humidity (%).", "2. ")
    add_bullet("Recharts Multi-Area Telemetry Plot: Dual Y-axes with gradient shading, custom hover tooltips, and highlighted anomaly window (ReferenceArea).", "3. ")

    add_h2("3.4 Anomaly Diagnostics & Explainability (WhyFlagged.tsx)")
    doc.add_paragraph("Explainable AI (XAI) cockpit answering why a specific anomaly was flagged:")
    add_bullet("SHAP Feature Attribution Bar Charts: Displays exact parameter contribution weights with direction (increases_anomaly vs decreases_anomaly).", "1. ")
    add_bullet("Multi-Model Confidence Matrix: Detailed score breakdown across Statistical Z-Score, PyTorch LSTM Autoencoder, Isolation Forest, and Spatial Consensus.", "2. ")
    add_bullet("Domain-Grounded Meteorological Narrative: Automated natural language explanation synthesizing physical weather dynamics.", "3. ")

    add_h2("3.5 Sensor Health & Prognostics (SensorHealth.tsx)")
    doc.add_paragraph("Predictive maintenance dashboard forecasting hardware reliability:")
    add_bullet("Animated SVG Circular Health Gauge: Dynamic 0–100 health score with smooth stroke-dashoffset transitions (>=80 Green, >=50 Amber, <50 Rose).", "1. ")
    add_bullet("Degradation Trajectory: Trend indicator (Improving, Stable, Degrading) based on historical drift and variance.", "2. ")
    add_bullet("Maintenance Timeline Forecast: Recommended maintenance inspection interval (e.g., 'Service recommended within 14 days').", "3. ")
    add_bullet("Subsystem Diagnostic Bars: Individual health meters for RTD Temperature Probe, Barometric Capsule, Hygrometer, and Telemetry Modem.", "4. ")

    add_h2("3.6 Alert History Archive (History.tsx)")
    doc.add_paragraph("Historical repository for auditing and regulatory review:")
    add_bullet("Archive Tabs: Acknowledged, Resolved, and All Historical records.", "1. ")
    add_bullet("Audit Timeline: Complete log of resolved anomalies with root-cause labels, timestamp conversion, and historical search.", "2. ")

    # -------------------------------------------------------------
    # 4. DESIGN SYSTEM & CSS TOKENS
    # -------------------------------------------------------------
    add_h1("4. Design System, CSS Variables & Tokens")
    doc.add_paragraph("All styling adheres to a modern, dark-mode glassmorphic design token system defined in frontend/src/index.css:")

    add_code_block(
        "/* Core Design Tokens */\n"
        "--bg-dark: #090d16;\n"
        "--bg-card: rgba(17, 24, 39, 0.85);\n"
        "--border-card: rgba(255, 255, 255, 0.08);\n"
        "--text-primary: #f8fafc;\n"
        "--text-secondary: #94a3b8;\n"
        "--accent-cyan: #38bdf8;\n"
        "--accent-emerald: #10b981;\n"
        "--accent-amber: #f59e0b;\n"
        "--accent-rose: #f43f5e;\n\n"
        "/* Glass Panel Utility */\n"
        ".glass-panel {\n"
        "  background: var(--bg-card);\n"
        "  backdrop-filter: blur(16px);\n"
        "  border: 1px solid var(--border-card);\n"
        "  border-radius: 12px;\n"
        "  box-shadow: 0 4px 20px -2px rgba(0, 0, 0, 0.5);\n"
        "}"
    )

    # -------------------------------------------------------------
    # 5. DATA CLIENT & STATE LIFECYCLE
    # -------------------------------------------------------------
    add_h1("5. Typed API Client & EventSource Architecture")
    doc.add_paragraph("The frontend utilizes a modular TypeScript API service (frontend/src/api.ts):")
    add_bullet("fetchHealth(): Polls GET /health to maintain navbar status badge.", "• ")
    add_bullet("fetchStations(): Queries GET /stations for map and station grid cards.", "• ")
    add_bullet("fetchAlerts(params): Queries GET /alerts with multi-field search/filter parameters.", "• ")
    add_bullet("fetchTimeseries(id, hours): Requests GET /stations/{id}/timeseries for Recharts plotting.", "• ")
    add_bullet("fetchSensorHealth(id): Retrieves GET /sensor-health/{id} for prognostic health scores.", "• ")
    add_bullet("fetchExplanation(alertId): Retrieves GET /explain/{alertId} for SHAP charts and narratives.", "• ")
    add_bullet("acknowledgeAlert(alertId): Sends POST /alerts/{alertId}/acknowledge for operator triage.", "• ")
    add_bullet("createAlertEventSource(onMessage): Establishes persistent SSE stream to /alerts/stream with auto-reconnection.", "• ")

    output_path = r"c:\SkyguardAI\SkyguardAI_Frontend_Documentation.docx"
    doc.save(output_path)
    print(f"Frontend Document successfully created at {output_path}")

if __name__ == "__main__":
    create_frontend_document()
