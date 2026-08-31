import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_techstack_document():
    doc = docx.Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    
    COLOR_PRIMARY = RGBColor(15, 32, 39)    # Deep Navy
    COLOR_ACCENT = RGBColor(0, 131, 176)   # Cyan / Teal
    COLOR_DARK = RGBColor(44, 62, 80)      # Charcoal Dark
    COLOR_MUTED = RGBColor(100, 110, 120)  # Gray
    
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = COLOR_DARK
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # HEADER / TITLE
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(15)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("🛰️ SkyguardAI")
    run_title.font.name = 'Segoe UI'
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Technology Stack & Implementation Methodology Documentation")
    run_sub.font.name = 'Segoe UI'
    run_sub.font.size = Pt(15)
    run_sub.font.bold = True
    run_sub.font.color.rgb = COLOR_ACCENT

    # Metadata box
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    metadata = [
        ("Platform:", "SkyguardAI Atmospheric Anomaly Detection & Consensus Platform"),
        ("Architecture:", "FastAPI + PyTorch LSTM Autoencoder + React 19 + SSE Streaming"),
        ("Deployment Targets:", "Automatic Weather Stations (AWS) & IMD Antarctic Maitri Network"),
        ("Document Version:", "1.0.0 (Production Architecture Specification)")
    ]
    
    idx = 0
    for r in range(2):
        for c in range(2):
            cell = meta_table.cell(r, c)
            set_cell_background(cell, "F4F7F6")
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            k, v = metadata[idx]
            run_k = p.add_run(f"{k} ")
            run_k.font.bold = True
            run_k.font.size = Pt(9.5)
            run_k.font.color.rgb = COLOR_PRIMARY
            run_v = p.add_run(v)
            run_v.font.size = Pt(9.5)
            run_v.font.color.rgb = COLOR_DARK
            idx += 1
            
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_section_header(title, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        if level == 1:
            p.paragraph_format.space_before = Pt(20)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(title)
            run.font.name = 'Segoe UI'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = COLOR_PRIMARY
        elif level == 2:
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(title)
            run.font.name = 'Segoe UI'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = COLOR_ACCENT
        elif level == 3:
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(title)
            run.font.name = 'Segoe UI'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = COLOR_DARK

    def create_styled_table(headers, rows, col_widths=None):
        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Header Row
        hdr_cells = table.rows[0].cells
        for i, header_text in enumerate(headers):
            hdr_cells[i].text = header_text
            set_cell_background(hdr_cells[i], "0F2027")
            set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = 'Segoe UI'
                run.font.size = Pt(9.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

        # Body Rows
        for r_idx, row_data in enumerate(rows):
            row_cells = table.rows[r_idx + 1].cells
            bg_color = "FFFFFF" if r_idx % 2 == 0 else "F9FBFC"
            for c_idx, cell_value in enumerate(row_data):
                row_cells[c_idx].text = str(cell_value)
                set_cell_background(row_cells[c_idx], bg_color)
                set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=140, right=140)
                p = row_cells[c_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = 'Segoe UI'
                    run.font.size = Pt(9)
                    run.font.color.rgb = COLOR_DARK

        if col_widths:
            for row in table.rows:
                for c_idx, width in enumerate(col_widths):
                    row.cells[c_idx].width = Inches(width)

        doc.add_paragraph().paragraph_format.space_after = Pt(8)
        return table

    # -------------------------------------------------------------
    # 1. TECHNOLOGIES USED
    # -------------------------------------------------------------
    add_section_header("1. Technologies Used & Architecture Stack", level=1)
    
    p = doc.add_paragraph()
    p.add_run("SkyguardAI is built upon a high-performance modern tech stack spanning asynchronous microservices, deep learning sequence modeling, spatial statistics, Explainable AI (XAI), real-time Server-Sent Events (SSE) streaming, and an operational Glassmorphism command dashboard.")

    add_section_header("1.1 Programming Languages & Runtime", level=2)
    create_styled_table(
        ["Language / Runtime", "Version", "Architectural Role"],
        [
            ["Python", "3.11+", "Backend server, PyTorch neural networks, spatial physics engines, and feature pipelines."],
            ["TypeScript", "5.8+", "Strict client-side typing, SSE event bus handling, and UI state architecture."],
            ["HTML5 / CSS3", "Modern Standard", "Semantic DOM layouts, Glassmorphism design tokens, CSS variables & animations."],
            ["SQL / Parquet", "Parquet 2.0", "High-throughput columnar storage for historical and simulated weather telemetry."]
        ],
        [1.8, 1.0, 3.7]
    )

    add_section_header("1.2 Backend Frameworks & Real-Time Engine", level=2)
    create_styled_table(
        ["Technology", "Version", "Description & Operational Utility"],
        [
            ["FastAPI", "0.110.0+", "Asynchronous ASGI web framework with automatic OpenAPI docs and dependency injection."],
            ["Uvicorn", "0.28.0+", "High-performance ASGI server with uvloop event loop and lightning-fast request routing."],
            ["SSE-Starlette", "2.0.0+", "Server-Sent Events streaming engine broadcasting live alerts without client polling."],
            ["Pydantic", "2.6.0+", "Strict data validation and domain Enum enforcement (Severity, RootCause, StationStatus)."],
            ["HTTPX", "0.27.0+", "Async HTTP client for backend smoke testing and integration verification."]
        ],
        [1.8, 1.0, 3.7]
    )

    add_section_header("1.3 Machine Learning, AI & Scientific Libraries", level=2)
    create_styled_table(
        ["Library", "Version", "Role in Anomaly Detection & Prognostics"],
        [
            ["PyTorch", "2.2.0+", "2-layer LSTM Autoencoder (64->32->64) detecting temporal sequence anomalies via MSE."],
            ["Scikit-Learn", "1.4.0+", "Multivariate Isolation Forest, standard scaling, and precision/recall evaluation metrics."],
            ["SHAP", "0.44.0+", "Explainable AI computing exact Shapley feature attributions (phi values) per alert."],
            ["Statsmodels", "0.14.0+", "Seasonal-Trend decomposition (STL) isolating 24h diurnal cycles from anomalies."],
            ["SciPy", "1.12.0+", "Mahalanobis distance calculation with covariance matrix for thermodynamic consistency."],
            ["Pandas & NumPy", "2.2+ / 1.26+", "Vectorized math, spatial Haversine distance, rolling averages, and time differentials."],
            ["PyArrow / Fastparquet", "15.0+ / 2024.2+", "Fast columnar serialization for multi-station telemetry datasets."],
            ["netCDF4 & xarray", "1.6.5+ / 2024.2+", "Decodes multi-dimensional scientific climate files from global meteorological centers."]
        ],
        [1.8, 1.0, 3.7]
    )

    add_section_header("1.4 Frontend Ecosystem & Visual Analytics", level=2)
    create_styled_table(
        ["Framework / Tool", "Version", "Purpose in Operational Dashboard"],
        [
            ["React", "19.2.8", "Component-driven reactive UI architecture with concurrent rendering and hooks."],
            ["Vite", "8.2.2", "Ultra-fast ESM build tool with instant Hot Module Replacement (HMR)."],
            ["React Router DOM", "6.22.0", "Client-side SPA routing across 6 core operational views."],
            ["Leaflet & React-Leaflet", "1.9.4 / 5.0.0", "Interactive CartoDB DarkMatter geospatial map with dynamic status pins."],
            ["Recharts", "3.10.1", "Declarative SVG charting for 72h telemetry curves, SHAP bars, and health gauges."],
            ["Lucide React", "1.37.0", "Meteorological and diagnostic SVG icons with clean, modern aesthetics."],
            ["Custom Glassmorphism CSS", "Native", "Futuristic dark operations theme (#070d18), glowing cards, and status rings."]
        ],
        [1.8, 1.0, 3.7]
    )

    add_section_header("1.5 Hardware, Sensors & AWS Weather Station Integration", level=2)
    create_styled_table(
        ["Sensor / Hardware", "Transducer Specification", "Operational Range", "Accuracy / Resolution"],
        [
            ["Ambient Temperature", "PT100 4-Wire RTD Platinum / Thermistor", "-60°C to +50°C", "±0.1°C / 0.01°C"],
            ["Barometric Pressure", "Piezoresistive Silicon Sensor", "750 to 1080 hPa", "±0.1 hPa / 0.05 hPa"],
            ["Relative Humidity", "Capacitive Thin-Film Polymer Hygrometer", "0% to 100% RH", "±1.5% RH / 0.1%"],
            ["Wind Speed & Direction", "Ultrasonic 2D/3D & Anemometer Vane", "0 to 200 knots", "±1 knot / ±2°"],
            ["Dataloggers / Compute", "Campbell CR1000X / ESP32-S3 / RPi CM4", "-40°C to +70°C", "16-bit ADC, RS-485 Modbus"],
            ["Field Telemetry", "4G LTE-M / NB-IoT / Iridium Satellite SBD", "Global Coverage", "15-min to 1-min Burst"],
            ["Polar Research Ground Truth", "IMD Antarctic Maitri Station (-70.75°S, 11.74°E)", "Extreme Antarctic", "Continuous Multi-Decadal Ground Truth"]
        ],
        [1.8, 2.0, 1.4, 1.3]
    )

    # -------------------------------------------------------------
    # 2. METHODOLOGY AND PROCESS FOR IMPLEMENTATION
    # -------------------------------------------------------------
    add_section_header("2. Methodology & Implementation Process", level=1)
    
    p = doc.add_paragraph()
    p.add_run("The SkyguardAI implementation follows a disciplined 5-stage pipeline designed for mission-critical environmental intelligence, strict false-positive suppression, and operator transparency.")

    add_section_header("Stage 1: Ingestion, Physical QC & Feature Engineering", level=2)
    p = doc.add_paragraph()
    p.add_run("Raw telemetry from physical AWS dataloggers or simulated networks is ingested and validated against physical boundary thresholds. The pipeline computes 1h, 3h, and 6h time differentials, rolling statistics (6h, 24h moving mean and standard deviation), cyclical trigonometric timestamps (sin/cos of hour and month), and LOESS STL decomposition to separate diurnal thermal cycles from anomalies.")

    add_section_header("Stage 2: Multi-Model Anomaly Detection Ensemble", level=2)
    p = doc.add_paragraph()
    p.add_run("Three distinct detection paradigms run concurrently on each observation window:\n"
              "• Statistical Detector: Evaluates rolling Z-scores (|Z| > 3.0) and frozen sensor zero-variance.\n"
              "• PyTorch LSTM Autoencoder: Computes reconstruction Mean Squared Error (MSE) across 24-step multi-sensor sequences.\n"
              "• Isolation Forest: Tree-partitioning algorithm isolating high-dimensional multivariate outliers.")

    add_section_header("Stage 3: Spatial Consensus & Thermodynamic Validation", level=2)
    p = doc.add_paragraph()
    p.add_run("The core innovation of SkyguardAI is the spatial and thermodynamic consensus engine. An anomaly flagged at a single station is cross-referenced with all neighboring stations within a 150km radius using Haversine great-circle distance. Elevation differences are normalized via physical atmospheric lapse rates (-6.5°C/km for temperature, -11 hPa/100m for pressure). Inter-parameter physical consistency is validated using Multivariate Mahalanobis distance. If neighboring stations exhibit the same deviation, the system classifies the event as a genuine regional weather front rather than a sensor defect.")

    add_section_header("Stage 4: Explainable AI (XAI) & Sensor Prognostics", level=2)
    p = doc.add_paragraph()
    p.add_run("For every detected alert, Tree/Kernel SHAP values quantify exact feature attributions, which are automatically translated into natural language meteorological narratives for human operators. Concurrently, a 0–100 Sensor Health Index is calculated with degradation trend classification (improving, stable, degrading) and Remaining Useful Life (RUL) maintenance forecasting.")

    add_section_header("Stage 5: Presentation & Real-Time Operational Dashboard", level=2)
    p = doc.add_paragraph()
    p.add_run("FastAPI and SSE-Starlette broadcast live alerts over a low-latency EventSource channel. The React 19 dashboard renders 6 dedicated operations views: Network Overview, Live Alerts with operator triage, 72h Station Detail telemetry graphs, Why Flagged XAI cockpit, Sensor Health prognostics, and Historical Audit Trail.")

    # -------------------------------------------------------------
    # 3. WORKING PROTOTYPE & EXECUTION GUIDE
    # -------------------------------------------------------------
    add_section_header("3. Working Prototype & Verification", level=1)
    
    p = doc.add_paragraph()
    p.add_run("The complete working prototype can be launched locally or containerized in Docker. Real-world fault scenarios can be dynamically injected to verify detection, spatial consensus, and XAI narratives in real-time.")

    add_section_header("3.1 Quickstart Commands", level=2)
    p = doc.add_paragraph()
    p.add_run("1. Backend API Server (FastAPI + Uvicorn):\n"
              "   .venv\\Scripts\\activate\n"
              "   pip install -r requirements.txt\n"
              "   uvicorn backend.app:app --host 0.0.0.0 --port 8000 --reload\n\n"
              "2. Frontend Operations Dashboard (React 19 + Vite):\n"
              "   cd frontend && npm install && npm run dev\n\n"
              "3. Docker Orchestration:\n"
              "   docker-compose up --build -d\n\n"
              "4. Automated Test Verification:\n"
              "   pytest tests/ -v")

    output_path = os.path.join(os.path.dirname(__file__), "SkyguardAI_TechStack_Documentation.docx")
    doc.save(output_path)
    print(f"Successfully generated styled Word document at: {output_path}")

if __name__ == "__main__":
    create_techstack_document()
